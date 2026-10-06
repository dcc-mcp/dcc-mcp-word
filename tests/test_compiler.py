"""Headless DOCX compiler tests — the `verified` half of the surface.

These run with no Word installation. That is the point: the compile path is
graded `verified` because it is provable headlessly, while reflow and
field-result refresh are `host_limited` (see test_capabilities).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from dcc_mcp_word.compiler import compile_document
from dcc_mcp_word.word_ir import parse_envelope

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def envelope(**overrides: object):
    base = {
        "schema_version": "office-ir/1.0",
        "kind": "word",
        "document_id": "draft:test",
        "metadata": {"title": "Test doc", "author": "pipeline", "language": "en"},
        "document": {},
        "outputs": ["docx"],
    }
    base.update(overrides)
    return parse_envelope(base)


def _field_instructions(path: Path) -> list[str]:
    document = DocxDocument(str(path))
    roots = [document.element.body]
    for section in document.sections:
        roots.append(section.header._element)
        roots.append(section.footer._element)
    found: list[str] = []
    for root in roots:
        for node in root.findall(f".//{_W_NS}instrText"):
            found.append((node.text or "").strip())
    return found


def test_compile_writes_paragraphs_and_sections(tmp_path: Path) -> None:
    env = envelope(
        document={
            "paragraphs": [{"text": "Intro", "style": "Normal"}],
            "sections": [{"title": "Summary", "paragraphs": [{"text": "Body", "style": "Normal"}]}],
        }
    )
    out = compile_document(env, tmp_path / "doc")
    assert out.suffix == ".docx"
    document = DocxDocument(str(out))
    texts = [p.text for p in document.paragraphs]
    assert "Intro" in texts
    assert "Summary" in texts
    assert "Body" in texts
    # A section title becomes a Heading 1 so the TOC can index it.
    assert document.paragraphs[texts.index("Summary")].style.name == "Heading 1"


def test_compile_writes_lists(tmp_path: Path) -> None:
    env = envelope(document={"lists": [{"items": ["alpha", "beta"], "style": "List Bullet"}]})
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    texts = [p.text for p in document.paragraphs]
    assert texts == ["alpha", "beta"]
    assert document.paragraphs[0].style.name == "List Bullet"


def test_compile_writes_numbered_list(tmp_path: Path) -> None:
    env = envelope(document={"lists": [{"items": ["one"], "style": "List Number"}]})
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    assert document.paragraphs[0].style.name == "List Number"


def test_compile_writes_tables(tmp_path: Path) -> None:
    env = envelope(document={"tables": [{"header": True, "rows": [["Shot", "Status"], ["sh010", "ip"]]}]})
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Shot", "Status"]
    assert [c.text for c in table.rows[1].cells] == ["sh010", "ip"]
    header_bold = [run.font.bold for p in table.rows[0].cells[0].paragraphs for run in p.runs]
    assert any(header_bold)


def test_compile_skips_empty_tables(tmp_path: Path) -> None:
    env = envelope(document={"tables": [{"header": True, "rows": []}]})
    out = compile_document(env, tmp_path / "doc")
    assert not DocxDocument(str(out)).tables


def test_compile_writes_headers_footers_and_page_field(tmp_path: Path) -> None:
    env = envelope(
        document={
            "headers": [{"section_index": 0, "text": "Header text"}],
            "footers": [{"section_index": 0, "text": "Footer text"}],
            "fields": [{"kind": "page"}],
        }
    )
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    # python-docx exposes _Header/_Footer (paragraphs only, no .text).
    header_text = "".join(p.text for p in document.sections[0].header.paragraphs)
    footer_text = "".join(p.text for p in document.sections[0].footer.paragraphs)
    assert header_text.strip() == "Header text"
    assert "Footer text" in footer_text
    # The PAGE field lives in the footer, not in the body.
    assert "PAGE" in _field_instructions(out)


def test_compile_writes_toc_field_instruction(tmp_path: Path) -> None:
    env = envelope(document={"fields": [{"kind": "toc"}]})
    out = compile_document(env, tmp_path / "doc")
    instructions = _field_instructions(out)
    assert any(i.startswith("TOC") for i in instructions)


def test_compile_writes_custom_field(tmp_path: Path) -> None:
    env = envelope(document={"fields": [{"kind": "custom", "code": "NUMPAGES"}]})
    out = compile_document(env, tmp_path / "doc")
    assert "NUMPAGES" in _field_instructions(out)


def test_compile_writes_metadata(tmp_path: Path) -> None:
    env = envelope(metadata={"title": "Call sheet", "author": "pipeline", "language": "en"})
    out = compile_document(env, tmp_path / "doc")
    properties = DocxDocument(str(out)).core_properties
    assert properties.title == "Call sheet"
    assert properties.author == "pipeline"


def test_compile_inserts_page_break_before_section(tmp_path: Path) -> None:
    env = envelope(document={"sections": [{"title": "A", "page_break_before": True, "paragraphs": []}]})
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    assert any("w:br" in p._p.xml and 'w:type="page"' in p._p.xml for p in document.paragraphs)


def test_compile_creates_sections_for_later_indices(tmp_path: Path) -> None:
    env = envelope(document={"footers": [{"section_index": 2, "text": "late"}]})
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    assert len(document.sections) >= 3
    footer_text = "".join(p.text for p in document.sections[2].footer.paragraphs)
    assert footer_text.strip() == "late"


def test_compile_falls_back_to_normal_for_unknown_style(tmp_path: Path) -> None:
    env = envelope(document={"paragraphs": [{"text": "x", "style": "No Such Style"}]})
    out = compile_document(env, tmp_path / "doc")
    document = DocxDocument(str(out))
    assert document.paragraphs[0].style.name == "Normal"


def test_compile_adds_docx_suffix_when_missing(tmp_path: Path) -> None:
    env = envelope(document={"paragraphs": [{"text": "x"}]})
    out = compile_document(env, tmp_path / "no-suffix")
    assert out.name == "no-suffix.docx"


def test_compile_summary_counts(tmp_path: Path) -> None:
    from dcc_mcp_word.compiler import WordCompiler

    env = envelope(
        document={
            "paragraphs": [{"text": "a"}],
            "lists": [{"items": ["b", "c"]}],
            "tables": [{"header": False, "rows": [["x"]]}],
            "sections": [{"title": "S", "paragraphs": [{"text": "d"}]}],
        }
    )
    compiler = WordCompiler(env)
    compiler.compile(tmp_path / "doc")
    assert compiler.summary["paragraphs"] == 2
    assert compiler.summary["lists"] == 2
    assert compiler.summary["tables"] == 1
    assert compiler.summary["sections"] == 1
