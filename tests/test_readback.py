"""Read-back verification tests — the 1.0 "write then read it back" gate."""

from __future__ import annotations

from pathlib import Path

from dcc_mcp_word.compiler import compile_document
from dcc_mcp_word.readback import read_back
from dcc_mcp_word.word_ir import parse_envelope


def envelope(**overrides: object):
    base = {
        "schema_version": "office-ir/1.0",
        "kind": "word",
        "document_id": "draft:test",
        "metadata": {"title": "Test", "author": "pipeline", "language": "en"},
        "document": {},
        "outputs": ["docx"],
    }
    base.update(overrides)
    return parse_envelope(base)


def test_read_back_passes_for_a_compiled_document(tmp_path: Path) -> None:
    env = envelope(
        document={
            "paragraphs": [{"text": "Intro", "style": "Normal"}],
            "lists": [{"items": ["a", "b"], "style": "List Bullet"}],
            "sections": [{"title": "Summary", "paragraphs": [{"text": "Body"}]}],
        }
    )
    out = compile_document(env, tmp_path / "doc.docx")
    report = read_back(env, out)
    assert report.ok, report.mismatches
    # Intro + 2 list items + the section heading + the section body.
    assert report.checked_paragraphs == 5
    assert report.mismatches == ()


def test_read_back_compares_table_cells(tmp_path: Path) -> None:
    env = envelope(document={"tables": [{"header": True, "rows": [["Shot", "Status"], ["sh010", "ip"]]}]})
    out = compile_document(env, tmp_path / "doc.docx")
    report = read_back(env, out)
    assert report.ok, report.mismatches
    assert report.checked_tables == 1
    assert report.checked_cells == 4


def test_read_back_detects_a_table_cell_mismatch(tmp_path: Path) -> None:
    env = envelope(document={"tables": [{"header": False, "rows": [["expected"]]}]})
    out = compile_document(env, tmp_path / "doc.docx")
    # Tamper with the artifact: the gate must catch it, not pass silently.
    from docx import Document as DocxDocument

    document = DocxDocument(str(out))
    document.tables[0].rows[0].cells[0].text = "tampered"
    document.save(str(out))

    report = read_back(env, out)
    assert not report.ok
    assert any(m.location.endswith("[0][0]") for m in report.mismatches)


def test_read_back_detects_a_paragraph_mismatch(tmp_path: Path) -> None:
    env = envelope(document={"paragraphs": [{"text": "original"}]})
    out = compile_document(env, tmp_path / "doc.docx")
    from docx import Document as DocxDocument

    document = DocxDocument(str(out))
    document.paragraphs[0].runs[0].text = "changed"
    document.save(str(out))

    report = read_back(env, out)
    assert not report.ok


def test_read_back_detects_a_paragraph_count_mismatch(tmp_path: Path) -> None:
    env = envelope(document={"paragraphs": [{"text": "one"}]})
    out = compile_document(env, tmp_path / "doc.docx")
    from docx import Document as DocxDocument

    document = DocxDocument(str(out))
    document.add_paragraph("extra")
    document.save(str(out))

    report = read_back(env, out)
    assert not report.ok
    assert any(m.location == "paragraphs.count" for m in report.mismatches)


def test_read_back_ignores_page_break_paragraphs(tmp_path: Path) -> None:
    """`add_page_break()` emits an empty paragraph the IR has no node for.

    A page break is formatting, not content: read-back must not count it as a
    paragraph mismatch, or every page-broken section fails the gate.
    """
    env = envelope(
        document={
            "sections": [
                {"title": "A", "paragraphs": [{"text": "first"}]},
                {"title": "B", "page_break_before": True, "paragraphs": [{"text": "second"}]},
            ]
        }
    )
    out = compile_document(env, tmp_path / "doc.docx")
    report = read_back(env, out)
    assert report.ok, report.mismatches
    assert report.checked_paragraphs == 4


def test_read_back_counts_fields_without_comparing_results(tmp_path: Path) -> None:
    """Field results are produced by Word, so the gate counts, not compares."""
    env = envelope(document={"fields": [{"kind": "toc"}, {"kind": "page"}]})
    out = compile_document(env, tmp_path / "doc.docx")
    report = read_back(env, out)
    assert report.ok, report.mismatches
    assert report.fields_present == 2


def test_read_back_reports_headers_and_footers(tmp_path: Path) -> None:
    env = envelope(
        document={
            "headers": [{"section_index": 0, "text": "H"}],
            "footers": [{"section_index": 0, "text": "F"}],
        }
    )
    out = compile_document(env, tmp_path / "doc.docx")
    report = read_back(env, out)
    assert report.headers == 1
    assert report.footers == 1


def test_read_back_raises_for_a_missing_artifact(tmp_path: Path) -> None:
    env = envelope(document={"paragraphs": [{"text": "x"}]})
    with __import__("pytest").raises(FileNotFoundError):
        read_back(env, tmp_path / "absent.docx")


def test_read_back_report_serialises(tmp_path: Path) -> None:
    env = envelope(document={"paragraphs": [{"text": "x"}]})
    out = compile_document(env, tmp_path / "doc.docx")
    body = read_back(env, out).to_dict()
    assert body["ok"] is True
    assert body["path"] == str(out)
    assert "checked_paragraphs" in body
