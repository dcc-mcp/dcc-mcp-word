"""word-document / inspect_document — read-only inventory of an existing DOCX.

Parameter resolution order (dcc-mcp-core execute_script convention):
1. stdin JSON: {"input": ...}
2. CLI flags: --input
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document as DocxDocument

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _force_utf8_stdio() -> None:
    """Deterministic output contract: stdout/stderr are always UTF-8."""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def _count_fields(document: DocxDocument) -> int:
    roots = [document.element.body]
    for section in document.sections:
        roots.append(section.header._element)
        roots.append(section.footer._element)
    return sum(len(root.findall(f".//{_W_NS}instrText")) for root in roots)


def run(params: dict) -> None:
    path = Path(params["input"])
    if not path.is_file():
        print(json.dumps({"success": False, "message": f"document not found: {path}", "context": {}}, ensure_ascii=False))
        return

    document = DocxDocument(str(path))
    paragraphs = [
        {"style": p.style.name, "text": p.text}
        for p in document.paragraphs
    ]
    tables = [
        {
            "rows": len(table.rows),
            "columns": len(table.columns) if table.rows else 0,
            "cells": [[cell.text for cell in row.cells] for row in table.rows],
        }
        for table in document.tables
    ]
    headers = ["".join(p.text for p in s.header.paragraphs) for s in document.sections]
    footers = ["".join(p.text for p in s.footer.paragraphs) for s in document.sections]

    print(
        json.dumps(
            {
                "success": True,
                "message": f"inspected {path.name}: {len(paragraphs)} paragraph(s), {len(tables)} table(s)",
                "context": {
                    "path": str(path),
                    "title": document.core_properties.title,
                    "sections": len(document.sections),
                    "paragraph_count": len(paragraphs),
                    "table_count": len(tables),
                    "fields": _count_fields(document),
                    "headers": headers,
                    "footers": footers,
                    "paragraphs": paragraphs,
                    "tables": tables,
                },
            },
            ensure_ascii=False,
        )
    )


def main() -> None:
    _force_utf8_stdio()
    params: dict = {}
    if not sys.stdin.isatty():
        raw = sys.stdin.read()
        if raw.strip():
            try:
                params = json.loads(raw)
            except json.JSONDecodeError:
                params = {}
    if not params:
        parser = argparse.ArgumentParser(description="Read-only inventory of a DOCX document")
        parser.add_argument("--input", required=True, help="DOCX file path")
        params = vars(parser.parse_args())
    try:
        run(params)
    except Exception as exc:  # noqa: BLE001 — surface as structured error
        print(json.dumps({"success": False, "message": str(exc), "context": {}}, ensure_ascii=False))
        sys.exit(1)


if __name__ == "__main__":
    main()
