"""Read-back verification — the "write then read it back" 1.0 gate.

Compiling a DOCX proves a file was written; it does not prove the file
contains what the IR asked for. This module reopens the artifact with
python-docx and compares its structure against the envelope paragraph by
paragraph and table by table.

Fields are counted, not compared: a `TOC`/`PAGE` field has no result until
Word opens the document and refreshes it, which is precisely why
`word.toc.rebuild` and `word.fields.update` are graded `host_limited`
(see `capabilities`).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument

from .word_ir import WordEnvelope

MISMATCH_LIMIT = 20


@dataclass(frozen=True)
class Mismatch:
    location: str
    expected: Any
    actual: Any


@dataclass(frozen=True)
class ReadbackReport:
    path: str
    checked_paragraphs: int
    checked_tables: int
    checked_cells: int
    mismatches: tuple[Mismatch, ...]
    fields_present: int
    headers: int
    footers: int
    sections: int

    @property
    def ok(self) -> bool:
        return not self.mismatches

    def to_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "path": self.path,
            "checked_paragraphs": self.checked_paragraphs,
            "checked_tables": self.checked_tables,
            "checked_cells": self.checked_cells,
            "mismatches": [
                {"location": m.location, "expected": m.expected, "actual": m.actual} for m in self.mismatches
            ],
            "fields_present": self.fields_present,
            "headers": self.headers,
            "footers": self.footers,
            "sections": self.sections,
        }


def _is_page_break(paragraph: Any) -> bool:
    """True for the empty paragraph `add_page_break()` emits.

    A page break is formatting, not content: it is a real body paragraph in
    the XML but carries no text. Read-back skips these so the paragraph
    counts line up with the IR, which has no page-break node of its own.
    """
    if paragraph.text.strip():
        return False
    return 'w:type="page"' in paragraph._p.xml


def _expected_paragraph_texts(envelope: WordEnvelope) -> list[tuple[str, str]]:
    """Flatten every IR paragraph into (style, text) in document order."""
    document = envelope.document
    out: list[tuple[str, str]] = []
    for paragraph in document.paragraphs:
        out.append((paragraph.style or "Normal", paragraph.text))
    for list_block in document.lists:
        for item in list_block.items:
            out.append((list_block.style or "List Bullet", item))
    for section in document.sections:
        if section.title is not None:
            out.append(("Heading 1", section.title))
        for paragraph in section.paragraphs:
            out.append((paragraph.style or "Normal", paragraph.text))
    return out


def _count_fields(document: DocxDocument) -> int:
    """Count field instructions in the body, headers and footers."""
    count = 0
    roots = [document.element.body]
    for section in document.sections:
        for part in (section.header, section.footer):
            roots.append(part._element)
    for root in roots:
        count += len(root.findall(".//" + "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instrText"))
    return count


def _part_text(part: Any) -> str:
    """Text of a header/footer part.

    python-docx exposes `_Header`/`_Footer`, which have paragraphs but no
    `.text` attribute — join the paragraph texts instead.
    """
    try:
        paragraphs = part.paragraphs
    except AttributeError:
        return ""
    return "".join(p.text for p in paragraphs).strip()


def read_back(envelope: WordEnvelope, docx_path: str | Path) -> ReadbackReport:
    """Reopen `docx_path` and compare it against `envelope`."""
    path = Path(docx_path)
    if not path.is_file():
        raise FileNotFoundError(f"document artifact not found: {path}")
    document = DocxDocument(str(path))

    expected = _expected_paragraph_texts(envelope)
    # Skip formatting-only paragraphs: a TOC field occupies a body paragraph
    # of its own (verified through the field count, not as document text), and
    # `add_page_break()` emits an empty one that the IR has no node for.
    toc_count = sum(1 for f in envelope.document.fields if f.kind == "toc")
    body = [p for p in document.paragraphs if not _is_page_break(p)]
    if toc_count:
        body = body[toc_count:]
    actual = [(p.style.name, p.text) for p in body]

    mismatches: list[Mismatch] = []
    checked_paragraphs = 0
    for index, (expected_item, actual_item) in enumerate(zip(expected, actual)):
        checked_paragraphs += 1
        if expected_item != actual_item:
            mismatches.append(
                Mismatch(
                    location=f"paragraphs[{index}]",
                    expected={"style": expected_item[0], "text": expected_item[1]},
                    actual={"style": actual_item[0], "text": actual_item[1]},
                )
            )
            if len(mismatches) >= MISMATCH_LIMIT:
                break
    if len(expected) != len(actual):
        mismatches.append(
            Mismatch(
                location="paragraphs.count",
                expected=len(expected),
                actual=len(actual),
            )
        )

    checked_tables = 0
    checked_cells = 0
    expected_tables = [t for t in envelope.document.tables if t.rows]
    if len(expected_tables) != len(document.tables):
        mismatches.append(Mismatch(location="tables.count", expected=len(expected_tables), actual=len(document.tables)))
    for table_index, (table_ir, table) in enumerate(zip(expected_tables, document.tables)):
        checked_tables += 1
        for row_index, row in enumerate(table_ir.rows):
            if row_index >= len(table.rows):
                mismatches.append(Mismatch(location=f"tables[{table_index}].rows", expected=len(row), actual="missing"))
                break
            for column_index, value in enumerate(row):
                checked_cells += 1
                if column_index >= len(table.rows[row_index].cells):
                    mismatches.append(
                        Mismatch(location=f"tables[{table_index}][{row_index}][{column_index}]", expected=value, actual="missing")
                    )
                    continue
                actual_value = table.rows[row_index].cells[column_index].text
                if value.strip() != actual_value.strip():
                    mismatches.append(
                        Mismatch(location=f"tables[{table_index}][{row_index}][{column_index}]", expected=value, actual=actual_value)
                    )
                if len(mismatches) >= MISMATCH_LIMIT:
                    break
            if len(mismatches) >= MISMATCH_LIMIT:
                break
        if len(mismatches) >= MISMATCH_LIMIT:
            break

    headers = sum(1 for s in document.sections if _part_text(s.header))
    footers = sum(1 for s in document.sections if _part_text(s.footer))

    return ReadbackReport(
        path=str(path),
        checked_paragraphs=checked_paragraphs,
        checked_tables=checked_tables,
        checked_cells=checked_cells,
        mismatches=tuple(mismatches),
        fields_present=_count_fields(document),
        headers=headers,
        footers=footers,
        sections=len(document.sections),
    )
